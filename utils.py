import re
import spacy
from sentence_transformers import SentenceTransformer, util
import torch
from transformers import BertTokenizer, BertForSequenceClassification, T5Tokenizer, T5ForConditionalGeneration
import sqlite3
import docx
from skills_pipeline import extract_skills_from_sections
from pdfminer.pdfparser import PDFSyntaxError
from pdfminer.high_level import extract_text as extract_pdf_text
from datetime import datetime
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from io import BytesIO
import joblib
import logging
from langdetect import detect

# Set logging for pdfminer
logging.getLogger('pdfminer').setLevel(logging.WARNING)

# Configure logger for this module
logger = logging.getLogger(__name__)

# Load spaCy and Sentence-BERT
nlp = spacy.load("en_core_web_sm")
sbert_model = SentenceTransformer('all-MiniLM-L6-v2')

# Load BERT model
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
bert_model_path = "resume_bert_model"
bert_tokenizer = BertTokenizer.from_pretrained(bert_model_path)
bert_model = BertForSequenceClassification.from_pretrained(bert_model_path)
bert_model.to(device)
bert_model.eval()

# Load MultiLabelBinarizer
mlb = joblib.load("label_binarizer.pkl")  # Assumes saved binarizer

# Load T5 model
t5_tokenizer = T5Tokenizer.from_pretrained("t5-small")
t5_model = T5ForConditionalGeneration.from_pretrained("t5-small")
t5_model.to(device)


# --- TEXT EXTRACTION ---
def extract_docx_text(file_path):
    doc = docx.Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_text(file_path):
    ext = file_path.lower().split('.')[-1]
    text = ""

    try:
        if ext == 'pdf':
            text = extract_pdf_text(file_path)
        elif ext in ['doc', 'docx']:
            text = extract_docx_text(file_path)
        else:
            raise ValueError("Unsupported file type.")
    except PDFSyntaxError:
        raise PDFSyntaxError("The uploaded file is not a valid PDF.")
    except Exception as e:
        raise RuntimeError(f"Failed to extract text: {str(e)}")

    # Normalize newlines and whitespace
    text = text.replace('\r', '\n')
    text = re.sub(r'\n+', '\n', text)
    return text.strip()

# Skill and certification lists
known_skills = [
    "python", "java", "c++", "sql", "html", "css", "javascript", "react", "angular", "node","php",
    "aws","mysql", "azure", "linux", "git", "github", "tensorflow", "keras", "pytorch", "pandas",
    "numpy", "scikit-learn", "flask", "django", "agile", "scrum", "rest", "api",
    "machine learning", "data analysis", "cloud computing"
]
certifications = [
    "aws certified", "microsoft certified", "scrum master", "pmp", "oracle certified",
    "google cloud certified", "cissp", "ccna"
]



def clean_skills_section(text):
    text = re.sub(r'[^\w\s,-]', '', text.lower()).strip()
    return ' '.join(
        [word for word in text.split() if word in known_skills or any(cert in word for cert in certifications)])


section_headers = {
    "summary": ["summary", "professional summary", "about me"],
    "skills": ["skills", "technical skills", "core competencies"],
    "experience": ["experience", "work experience", "professional experience"],
    "education": ["education", "academic background", "qualifications"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "licenses", "certificates", "certification"],
    "languages": ["languages known", "languages"],
    "activities": ["extra-curricular", "activities", "achievements"]
}

SECTION_TITLES = [
    "summary", "objective", "skills", "experience", "education", "projects",
    "certifications", "certification", "certificates", "awards",
    "publications", "achievements"
]
SECTION_ALIASES = {
    "certification": "certifications",
    "certificates": "certifications"
}

def split_into_sections(text):
    """
    Split resume text into sections (list of lines per section).
    Detects section headers flexibly and groups following lines.
    """
    sections = {}
    current = None

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue

        lower = stripped.lower()
        normalized = re.sub(r'[^a-z]', '', lower)

        matched = False
        for title in SECTION_TITLES:
            norm_title = re.sub(r'[^a-z]', '', title.lower())
            if normalized == norm_title or normalized.startswith(norm_title):
                mapped = SECTION_ALIASES.get(title.lower(), title.lower())
                current = mapped
                sections[current] = []
                matched = True
                break

        if not matched and current:
            sections[current].append(stripped)

    # cleanup step – turn blobs into lists of items
    def clean_list(values, invalid_keywords=None):
        if not values:
            return []
        if invalid_keywords is None:
            invalid_keywords = []
        cleaned = []
        for v in values:
            txt = str(v).strip()
            if not txt:
                continue
            if any(kw.lower() in txt.lower() for kw in invalid_keywords):
                continue
            cleaned.append(txt)
        return cleaned

    # apply per-section cleaning
    cleaned_sections = {}
    for k, v in sections.items():
        if k == "skills":
            # flat list only
            skill_text = " ".join(v)
            parts = re.split(r"[,\|]", skill_text)
            skills = []
            for s in parts:
                s = re.sub(r"^[\u2022\uf0b7\-•\s]+", "", s).strip()
                if not s or len(s.split()) > 5:
                    continue
                if ":" in s:
                    s = s.split(":", 1)[1].strip()  # keep only actual skill
                skills.append(s)
            cleaned_sections[k] = list(dict.fromkeys(skills))

        elif k == "certifications":
            # clean course names, drop "CERTIFICATE"
            certs = []
            for c in v:
                c = re.sub(r"\|?\s*CERTIFICATE.*", "", c, flags=re.I).strip()
                c = re.sub(r"\s*Certified\s*", "", c, flags=re.I).strip()
                if c:
                    certs.append(c)
            cleaned_sections[k] = certs

        elif k == "projects":
            # keep all project lines (short titles only)
            projects = []
            for line in v:
                if len(line.split()) <= 15 and not re.search(r"(CGPA|%|Technologies Used|DOI)", line, re.I):
                    line = re.sub(r"\|.*", "", line)  # drop links
                    projects.append(line.strip())
            cleaned_sections[k] = projects

        else:
            cleaned_sections[k] = clean_list(v)

    return cleaned_sections

def detect_resume_sections(text: str):
    structured = {
        "summary": "",
        "skills": [],
        "projects": [],
        "experience": [],
        "certifications": [],
        "education": [],
        "languages": [],
        "activities": [],
        "publications": []
    }

    # --- Summary ---
    summary_match = re.search(r"(Professional summary|Summary)[:\-]?\s*(.*?)(?=\n\S+:|\Z)", text, re.S | re.I)
    if summary_match:
        structured["summary"] = summary_match.group(2).strip()

    # --- Skills ---
    skills_block = re.search(
        r"(Technical Skills|Skills)[:\-]?\s*(.*?)(?=\n(?:Education|Projects|Certifications?|Experience|Languages|Extra-Curricular|$))",
        text, re.S | re.I)
    if skills_block:
        raw_skills = skills_block.group(2)
        skills = re.split(r"[,\n]", raw_skills)
        cleaned = []
        for s in skills:
            s = re.sub(r"^[\u2022\uf0b7\-•\s]+", "", s).strip()
            if not s or len(s.split()) > 4:
                continue
            if any(x in s.lower() for x in ["industry", "technical skills", "experience"]):
                continue
            if ":" in s:
                parts = [p.strip() for p in s.split(":") if p.strip()]
                cleaned.extend(parts[1:] if len(parts) > 1 else parts)
            else:
                cleaned.append(s)
        structured["skills"] = list(dict.fromkeys(cleaned))

    # --- Projects ---
    def clean_projects(lines):
        clean_titles = []
        for line in lines:
            # Skip junk lines (dates, CGPA, %, institutes, etc.)
            if re.search(r"(\d{4}|\d{1,2}\.\d+%|CGPA|Present|Acharya|University|College|School)", line, re.I):
                continue
            # Skip "Technologies Used" or filler
            if line.lower().startswith("technologies used"):
                continue
            # Skip long sentences/descriptions
            if len(line.split()) > 12 or line.endswith((".", ":", ";")):
                continue
            # Extract only clean project title
            title = re.split(r"[\|:]", line)[0].strip()
            if title and title not in clean_titles:
                clean_titles.append(title)
        return clean_titles

    proj_block = re.search(
        r"(Projects)[:\-]?\s*(.*?)(?=\n(?:Certifications?|Education|Research|Experience|Languages|Extra|$))",
        text, re.S | re.I
    )
    if proj_block:
        raw_projects = proj_block.group(2)
        lines = [p.strip("•- \t") for p in raw_projects.split("\n") if p.strip()]
        structured["projects"] = clean_projects(lines)

    # --- Certifications ---
    cert_block = re.search(
        r"(Certifications?)[:\-]?\s*(.*?)(?=\n(?:Education|Projects|Languages|Research|Experience|Extra|$))",
        text, re.S | re.I
    )
    if cert_block:
        raw = cert_block.group(2)
        certs = []
        for c in raw.split("\n"):
            c = c.strip("•- \t")
            if not c:
                continue

            c = re.sub(r"\s*\|?\s*CERTIFICATE.*", "", c, flags=re.I).strip()
            # Skip headers
            if c.lower() in ["s:", "certifications"]:
                continue
            # Ensure it looks like a course name
            if len(c.split()) < 2:
                continue
            certs.append(c)
        # Deduplicate while keeping order
        structured["certifications"] = list(dict.fromkeys(certs))

    # Education
    edu_block = re.search(r"(Education)[:\-]?\s*(.*?)(?=\n(?:Projects|Certifications?|Languages|Extra-Curricular|$))", text, re.S | re.I)
    if edu_block:
        raw = edu_block.group(2)
        lines = [e.strip("•- \t") for e in raw.split("\n") if e.strip()]
        structured["education"] = [l for l in lines if not l.lower().startswith("cgpa")]

    # Publications
    pub_block = re.search(
        r"(Research|Publications|Research Publications)[:\-]?\s*(.*?)(?=\n(?:Certifications?|Education|Languages|Activities|$))",
        text, re.S | re.I)
    if pub_block:
        raw = pub_block.group(2)
        pubs = [p.strip("•- \t") for p in raw.split("\n") if p.strip()]
        pubs = [p for p in pubs if "certificate" not in p.lower() and "publications" not in p.lower()]
        structured["publications"] = pubs

    # Languages
    lang_block = re.search(r"(Languages Known|Languages)[:\-]?\s*(.*?)(?=\n(?:Extra-Curricular|$))", text, re.S | re.I)
    if lang_block:
        raw = lang_block.group(2)
        structured["languages"] = [l.strip("•- \t") for l in raw.split("\n") if l.strip()]

    # Activities
    act_block = re.search(r"(Extra-Curricular Activities|Activities)[:\-]?\s*(.*?)(?=\Z)", text, re.S | re.I)
    if act_block:
        raw = act_block.group(2)
        structured["activities"] = [a.strip("•- \t") for a in raw.split("\n") if a.strip()]

    return structured

def get_t5_feedback(resume_text):
    sections = split_into_sections(resume_text)
    relevant_sections = []
    for key in ["summary", "experience", "projects", "education", "skills"]:
        if key in sections:
            relevant_sections.append(sections[key])

    cleaned_text = " ".join([sec for sec in relevant_sections if sec]).strip()
    cleaned_text = re.sub(r'\s+', ' ', cleaned_text).strip()[:500]

    prompt = (
        "Generate 2–4 concise, actionable ATS-focused suggestions to improve this resume. "
        "Focus on adding keywords, quantifiable metrics, and clear formatting. "
        "Avoid generic terms, questions, raw resume content, or prompts. Output only complete sentences: "
        f"{cleaned_text}"
    )

    try:
        inputs = t5_tokenizer(prompt, return_tensors="pt", truncation=True, padding=True).to(device)
        with torch.no_grad():
            outputs = t5_model.generate(
                inputs['input_ids'],
                max_length=150,
                num_beams=4,
                early_stopping=True,
                do_sample=True,
                top_p=0.9,
                no_repeat_ngram_size=2  # Prevent repetition
            )
        suggestion_text = t5_tokenizer.decode(outputs[0], skip_special_tokens=True)
        suggestions = [s.strip() for s in re.split(r'(?<=[.!?])\s+', suggestion_text) if
                       s.strip() and len(s.strip()) > 10]

        invalid_patterns = [
            r'^(what|how|why|can|describe|explain)\b.*[.!?]$',
            r'^\s*(summary|skills|education|experience)\b',
            r'^\s*[\w\s,-]+(,[\w\s,-]+)+$',
            r'^\s*\d+\s*$',
            r'^\s*[\d\.\s%-]+$',
            r'\b(cgpa|percentage|portfolio|create)\b'
        ]
        filtered_suggestions = [
            s for s in suggestions
            if (len(s) > 20 and len(s) < 200 and
                not any(re.match(p, s, re.IGNORECASE) for p in invalid_patterns) and
                s.endswith(('.', '!')))
        ]

        if not filtered_suggestions:
            filtered_suggestions = [
                "Incorporate specific technical skills like Python or AWS to enhance ATS keyword matching.",
                "Add measurable outcomes in the experience section, such as 'improved performance by 20%'.",
                "Use clear section headers like 'Skills' and 'Experience' to improve ATS parsing."
            ]

        return [{"type": "Insights", "message": f" {s}"} for s in filtered_suggestions[:4]]
    except Exception as e:
        logging.error(f"T5 feedback generation failed: {str(e)}")
        return [{"type": "Insights", "message": " Unable to generate ATS insights due to processing error."}]

def advanced_style_checks(text):
    issues = []
    passive_patterns = [
        r'\b(is|was|were|are|been|being)\s+\w+ed\b',
        r'\b(get|got|gets)\s+\w+ed\b'
    ]
    for pattern in passive_patterns:
        if re.search(pattern, text, re.IGNORECASE):
            issues.append("Avoid passive voice; use active verbs to clarify ownership of results.")

    repeated = re.findall(r'\b(\w+)\s+\1\b', text.lower())
    if repeated:
        issues.append(f"Repeated words detected (e.g., '{', '.join(set(repeated))}'). Try to vary phrasing.")

    paragraphs = [p for p in text.split('\n') if len(p.strip()) > 120]
    if len(paragraphs) > 0:
        issues.append("Some sections have long paragraphs. Break them into bullet points for better readability.")

    weak_phrases = ["responsible for", "worked on", "involved in", "participated in", "helped with"]
    for phrase in weak_phrases:
        if phrase in text.lower():
            issues.append(f"Phrase like '{phrase}' is generic. Use more specific action verbs (e.g., 'led', 'built').")
            break

    return [{"type": "Style", "message": f" {msg}"} for msg in issues]

#ats format checking
def check_ats_formatting(resume_text, parsed_sections=None):
    issues = []
    required_sections = ["summary", "skills", "experience", "education", "projects", "certifications"]
    parsed_keys = parsed_sections.keys() if parsed_sections else []

    #Normalize sections
    experience_text = normalize_section(parsed_sections.get("experience", "")).strip()
    summary_text = normalize_section(parsed_sections.get("summary", "")).lower()

    experience_word_count = len(experience_text.split())
    has_valid_experience = experience_text and experience_word_count > 5

    # Build missing list
    missing = []
    for sec in required_sections:
        if sec == "experience" and not has_valid_experience:
            missing.append(sec.capitalize())
        elif sec != "experience" and sec not in parsed_keys:
            missing.append(sec.capitalize())

    if missing:
        issues.append(f"Resume is missing: {', '.join(missing)}.")

    # Check bullet points only for valid experience
    if has_valid_experience:
        if not any(b in experience_text for b in ["-", "•", "*"]):
            issues.append("Experience section lacks bullet points.")

    # Summary keywords
    if summary_text and not any(k in summary_text for k in ["experienced", "results-driven", "proven track record"]):
        issues.append("Add impact keywords like 'experienced' or 'results-driven' to your summary.")

    # Special characters
    special_chars = ['#', '*', '🎯', '✅', '🔥', '👉', '➤']
    if any(char in resume_text for char in special_chars):
        issues.append("Minimize special characters (e.g., *, #, emojis) to ensure ATS readability.")

    return [{"type": "ATS Optimization", "message": f"ATS Tips: - {msg}"} for msg in issues]

def check_language(resume_text):
    try:
        lang = detect(resume_text)
        if lang != 'en':
            return [{"type": "ATS Optimization",
                     "message": " Use English for your resume, as most ATS systems expect English."}]
    except:
        pass
    return []

def suggest_resume_improvements(
    text,
    detected_skills=None,
    resume_sections=None,
    job_description=None,
    job_title=None,
    detailed=True
):
    """
    Suggest resume improvements with structured JSON output.
    Sections:
    1. Missing Core Skills
    2. Missing Preferred Skills
    3. Missing JD Keywords
    4. Experience Alignment
    5. Role Alignment
    6. ATS & General Feedback
    """

    sections_output = []
    jd_lower = (job_description or "").lower()
    resume_skills = detected_skills or []

    # Extract JD segments
    core_skills_text = jd_lower
    preferred_skills_text = ""
    if "preferred qualifications" in jd_lower:
        parts = jd_lower.split("preferred qualifications")
        core_skills_text = parts[0]
        preferred_skills_text = parts[1]

    jd_core_skills = extract_skills_from_sections({"skills": core_skills_text})
    jd_preferred_skills = extract_skills_from_sections({"skills": preferred_skills_text})

    # --- 1. Missing Core Skills ---
    missing_core = [s for s in jd_core_skills if s not in resume_skills]
    if missing_core:
        sections_output.append({
            "title": "Missing Core Skills (from JD “Qualifications”)",
            "cv_has": resume_skills,
            "jd_has": jd_core_skills,
            "missing": missing_core,
            "suggestion": f"Explicitly add skills like {', '.join(missing_core[:10])}."
        })

    # --- 2. Missing Preferred Skills ---
    missing_preferred = [s for s in jd_preferred_skills if s not in resume_skills]
    if missing_preferred:
        sections_output.append({
            "title": "Missing Preferred Skills (JD “Nice to Have”)",
            "missing": missing_preferred,
            "suggestion": "If you have exposure to these, mention them — even as “familiar with” or from academic projects."
        })

    #3. Missing JD Keywords

    stopwords = {"and", "with", "from", "into", "your", "ability", "including", "various",
                 "fast", "good", "preferred", "teams", "environment", "skills", "knowledge"}
    jd_keywords = set(re.findall(r'\b[A-Za-z]{4,}\b', jd_lower)) - stopwords
    resume_keywords = set(re.findall(r'\b[A-Za-z]{4,}\b', text.lower()))
    missing_keywords = jd_keywords - resume_keywords
    if missing_keywords:
        sections_output.append({
            "title": "Missing JD Keywords",
            "jd_keywords": list(missing_keywords)[:12],
            "suggestion": "Add these keywords in your Summary, Skills, or Projects to boost ATS match."
        })

    #4. Experience Alignment
    if "experience" in jd_lower:
        exp_match = re.search(r'(\d+)\+?\s?(year|yr)', jd_lower)
        if exp_match:
            required_years = int(exp_match.group(1))
            resume_exp_match = re.search(r'(\d+)\+?\s?(year|yr)', text.lower())
            if not resume_exp_match or int(resume_exp_match.group(1)) < required_years:
                sections_output.append({
                    "title": "Experience Alignment",
                    "suggestion": f"JD asks for {required_years}+ years of experience. Frame your academic projects and internships as professional experience to better align with the JD."
                })
        else:
            sections_output.append({
                "title": "Experience Alignment",
                "suggestion": "JD mentions experience. Frame your academic projects and internships as professional experience to better align with the JD."
            })

    # 5. Role Alignment
    if job_title:
        if job_title.lower() not in text.lower():
            # Detect degree
            degree_patterns = {
                "mca": "MCA",
                "mba": "MBA",
                "be": "B.E.",
                "b.e": "B.E.",
                "btech": "B.Tech",
                "b.tech": "B.Tech",
                "bachelor": "Bachelor's",
                "master": "Master's"
            }
            degree_found = "Graduate"
            for key, val in degree_patterns.items():
                if re.search(rf"\b{key}\b", text.lower()):
                    degree_found = val
                    break

            # Detect branch
            branch_patterns = {
                "computer science": "Computer Science",
                "cse": "Computer Science",
                "ise": "Information Science",
                "it": "Information Technology",
                "information technology": "Information Technology",
                "ece": "Electronics and Communication",
                "electronics": "Electronics",
                "mechanical": "Mechanical",
                "civil": "Civil",
                "mba": "Business Administration",
            }
            branch_found = None
            for key, val in branch_patterns.items():
                if re.search(rf"\b{key}\b", text.lower()):
                    branch_found = val
                    break

            # Pick top 3 relevant skills from intersection of CV & JD
            skills_list = [
                "python", "java", "c++", "c#", "sql", "mysql", "javascript", "html", "css",
                "react", "angular", "node", "django", "flask", "git", "aws", "azure"
            ]
            cv_skills = {skill for skill in skills_list if skill in text.lower()}
            jd_skills = {skill for skill in skills_list if skill in jd_lower}
            common_skills = list(cv_skills & jd_skills) or list(cv_skills)
            skills_str = ", ".join(
                skill.capitalize() for skill in common_skills[:3]) if common_skills else "key technologies"

            # Build summary suggestion
            if degree_found == "MCA":
                summary_line = f"{degree_found} student aspiring to start a career as a {job_title}, skilled in {skills_str}."
            elif branch_found:
                summary_line = f"{degree_found} student ({branch_found}) aspiring to start a career as a {job_title}, skilled in {skills_str}."
            else:
                summary_line = f"{degree_found} student aspiring to start a career as a {job_title}, skilled in {skills_str}."

            sections_output.append({
                "title": "Role Alignment",
                "suggestion": f"JD says “{job_title}” but your CV summary does not explicitly mention it. Update your summary to: '{summary_line}'"
            })

    # 6. ATS & Section Feedback
    inputs = bert_tokenizer(text, return_tensors="pt", truncation=True, padding=True).to(device)
    with torch.no_grad():
        outputs = bert_model(**inputs)
        probs = torch.sigmoid(outputs.logits).cpu().numpy()[0]

    threshold = 0.5
    label_to_tip = {
        "missing_metrics": "Add quantifiable achievements (e.g., 'reduced processing time by 30%').",
        "format_issue": "Ensure consistent formatting — align headers, spacing, and bullet points.",
        "generic_summary": "Make your summary more specific — highlight major accomplishments and goals."
    }
    predicted_labels = [mlb.classes_[i] for i, p in enumerate(probs) if p > threshold]

    ats_tips = [label_to_tip[l] for l in predicted_labels if l in label_to_tip]
    if ats_tips:
        sections_output.append({
            "title": "ATS & General Resume Feedback",
            "tips": ats_tips
        })

    return {
        "labels": predicted_labels,
        "sections": sections_output,
        "suggestions": [
            sec.get("suggestion", "") if isinstance(sec.get("suggestion", ""), str)
            else " | ".join(sec.get("tips", []))
            for sec in sections_output
        ]
    }

def evaluate_section(name, text, skills=None, education=None):
    tips = []
    text = normalize_section(text)
    skills = skills or []

    if name == 'summary':
        if len(text.split()) < 30:
            tips.append("Expand your summary with specific achievements (e.g., 'Developed a Python-based web app').")
        if any(word in text.lower() for word in ['hardworking', 'team player', 'passionate']):
            tips.append(
                "Replace generic terms like 'hardworking' with ATS-friendly terms like 'developed' or 'implemented'.")

    if name == 'skills':
        if len(skills) < 5:
            tips.append("List at least 5–10 technical skills or certifications to improve ATS keyword matching.")
        if not any(skill in skills for skill in ['aws', 'pytorch', 'react']):
            tips.append("Add modern tools like AWS, PyTorch, or React to align with ATS expectations.")

    if name == 'experience':
        if not re.search(r'\d+%|\$|increased|reduced', text, re.IGNORECASE):
            tips.append("Include measurable results (e.g., 'increased efficiency by 15%') for ATS optimization.")
        if not re.search(r'[\*\-\•]\s', text):
            tips.append("Use bullet points to list responsibilities for better ATS parsing.")

    if name == 'education':
        if not re.search(r'\b\d{4}\b', text):
            tips.append("Add graduation year (e.g., 'MCA, 2024') to show recency for ATS systems.")
        if "mca" in text.lower():
            tips.append("Highlight MCA coursework like 'database systems' or 'algorithms' for ATS relevance.")

    return tips

def score_section(name, text):
    score = 0
    feedback = []
    text = normalize_section(text)

    found_skills = extract_skills_from_sections({"skills": text})
    if found_skills:
        score += min(len(found_skills) * 3, 15)
    else:
        feedback.append(f"{name.title()} section lacks technical skills or certifications.")

    if re.search(r'\d+%|\$|increased|reduced|saved|grew|improved', text, re.IGNORECASE):
        score += 10
    else:
        feedback.append(f"{name.title()} section lacks quantifiable metrics (e.g., 'reduced costs by 20%').")

    sentences = [s for s in text.split('.') if s.strip()]
    avg_len = sum(len(s.split()) for s in sentences) / max(1, len(sentences))
    if 10 <= avg_len <= 20:
        score += 10
    else:
        feedback.append(f"{name.title()} section has overly long or short sentences. Aim for 10–20 words per sentence.")

    if len(text.strip()) > 50:
        score += 5
    else:
        feedback.append(f"{name.title()} section is too brief. Add more details or achievements.")

    return min(score, 40), feedback

def extract_skills_from_text(text: str):
    """Backward-compatible wrapper"""
    return extract_skills_from_sections({"skills": text})


def normalize_section(section):
    """Ensure section is always a single string (not list/None)."""
    if isinstance(section, list):
        return " ".join(str(item) for item in section if item)
    elif section is None:
        return ""
    return str(section)

def score_resume_sections(sections):
    scores = {
        'summary': 0,
        'skills': 0,
        'experience': 0,
        'education': 0,
        'projects': 0
    }
    feedback = []

    # Normalize all sections
    summary = normalize_section(sections.get('summary', ''))
    skills_text = normalize_section(sections.get('skills', ''))
    projects_text = normalize_section(sections.get('projects', ''))
    experience_text = normalize_section(sections.get('experience', ''))
    education_text = normalize_section(sections.get('education', ''))

    #SUMMARY
    if summary:
        if len(summary.split()) > 20:
            scores['summary'] = 15
            if any(word in summary.lower() for word in ['experienced', 'proven', 'results']):
                scores['summary'] += 5
            else:
                feedback.append({
                    "type": "Summary",
                    "message": "Add impact keywords like 'experienced' or 'results-driven' to your summary."
                })
        else:
            feedback.append({
                "type": "Summary",
                "message": "Expand your summary to 2–3 impactful sentences."
            })
    else:
        feedback.append({
            "type": "Suggestion",
            "message": " Add a 'Summary' section to briefly introduce your background and strengths."
        })

    # SKILLS
    important_tech_skills = ['python', 'sql', 'react', 'aws', 'docker', 'firebase']
    if skills_text:
        found_skills = extract_skills_from_sections({"skills": skills_text})
        scores['skills'] = min(len(found_skills) * 4, 30)

        if len(found_skills) < 5:
            missing_skills = [s for s in important_tech_skills if s not in found_skills]
            if missing_skills:
                feedback.append({
                    "type": "Suggestion",
                    "message": f" Include important technical skills like {', '.join(missing_skills[:4])}."
                })
            else:
                feedback.append({
                    "type": "Skills",
                    "message": "List at least 5–10 technical skills or certifications."
                })
    else:
        feedback.append({
            "type": "Suggestion",
            "message": " Add a 'Skills' section with key technologies like Python, SQL, Firebase, etc."
        })

    # EXPERIENCE
    if experience_text:
        if any(word in experience_text.lower() for word in ['developed', 'led', 'built', 'managed']):
            scores['experience'] = 25
        else:
            scores['experience'] = 10
            feedback.append({
                "type": "Experience",
                "message": "Use action verbs like 'developed' or 'led' in your experience section."
            })
        if not re.search(r'\d+%|\$|increased|reduced', experience_text, re.IGNORECASE):
            feedback.append({
                "type": "Experience",
                "message": "Add quantifiable results (e.g., 'increased efficiency by 15%')."
            })
    else:
        feedback.append({
            "type": "Suggestion",
            "message": " Add an 'Experience' section to showcase internships, freelance work, or projects with real-world impact."
        })

    #EDUCATION
    if education_text:
        if any(word in education_text.lower() for word in ['bachelor', 'master', 'mca']):
            scores['education'] = 20
        else:
            scores['education'] = 10
            feedback.append({
                "type": "Education",
                "message": "Clearly mention your degree (e.g., 'MCA, 2024')."
            })
    else:
        feedback.append({
            "type": "Suggestion",
            "message": " Add an 'Education' section with your degree, university, and graduation year."
        })

    # PROJECTS
    if projects_text:
        if any(word in projects_text.lower() for word in ['developed', 'built', 'created']):
            scores['projects'] = 15
        else:
            scores['projects'] = 5
            feedback.append({
                "type": "Projects",
                "message": "Use action verbs like 'developed', 'built', or 'created' to describe your projects."
            })

        if not re.search(r'\d+%|\$|users|downloads|improved|faster', projects_text, re.IGNORECASE):
            feedback.append({
                "type": "Projects",
                "message": "Add measurable results (e.g., 'used by 100+ users' or 'reduced latency by 30%')."
            })
    else:
        feedback.append({
            "type": "Suggestion",
            "message": " Add a 'Projects' section to demonstrate practical applications of your skills."
        })

    total_score = sum(scores.values())
    return total_score, scores, feedback

def compute_similarity(resume_text, job_description):
    embeddings = sbert_model.encode([resume_text, job_description], convert_to_tensor=True)
    score = util.pytorch_cos_sim(embeddings[0], embeddings[1]).item()
    return round(score * 100, 2)


def init_db():
    conn = sqlite3.connect('results.db')
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS analysis_results (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        filename TEXT, option TEXT, score REAL, skills TEXT, feedback TEXT, timestamp TEXT
    )''')
    conn.commit()
    conn.close()


def save_result(filename, option, score, skills, feedback):
    conn = sqlite3.connect('results.db')
    c = conn.cursor()
    c.execute('''INSERT INTO analysis_results
                 (filename, option, score, skills, feedback, timestamp)
                 VALUES (?, ?, ?, ?, ?, ?)''',
              (filename, option, score, ", ".join(skills) if isinstance(skills, list) else skills,
               feedback, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    conn.commit()
    conn.close()

import unicodedata

def clean_text(text):
    if not isinstance(text, str):
        return ''
    return unicodedata.normalize('NFKD', text).encode('latin1', 'ignore').decode('latin1')


from io import BytesIO
from fpdf import FPDF
import os

def generate_pdf(data):
    pdf = FPDF()
    pdf.add_page()

    # ✅ Register DejaVuSans (Unicode safe)
    font_dir = os.path.join(os.path.dirname(__file__), "fonts")
    pdf.add_font('DejaVu', '', os.path.join(font_dir, 'DejaVuSans.ttf'), uni=True)
    pdf.add_font('DejaVu', 'B', os.path.join(font_dir, 'DejaVuSans-Bold.ttf'), uni=True)

    pdf.set_font("DejaVu", "B", 16)
    pdf.cell(0, 10, "Resume Analysis Report", ln=True)

    pdf.set_font("DejaVu", "", 12)
    pdf.ln(5)

    # Analysis Section
    analysis = data.get("analysisResults", {})
    if analysis:
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(0, 10, "Scores", ln=True)
        pdf.set_font("DejaVu", "", 12)
        pdf.cell(0, 10, f"Resume Score: {analysis.get('matchScore', 'N/A')}%", ln=True)
        pdf.cell(0, 10, f"Skills Identified: {analysis.get('skillsFound', 'N/A')}", ln=True)
        pdf.cell(0, 10, f"Recommendation: {analysis.get('recommendation', 'N/A')}", ln=True)
        pdf.ln(5)

        feedback_grouped = analysis.get("feedbackGrouped", {})
        if feedback_grouped:
            pdf.set_font("DejaVu", "B", 14)
            pdf.cell(0, 10, "Feedback", ln=True)
            pdf.set_font("DejaVu", "", 12)

            for category, messages in feedback_grouped.items():
                if not messages:
                    continue
                pdf.set_font("DejaVu", "B", 12)
                pdf.cell(0, 8, category, ln=True)
                pdf.set_font("DejaVu", "", 11)
                if len(messages) == 1:
                    pdf.multi_cell(0, 8, messages[0])
                else:
                    for msg in messages:
                        pdf.multi_cell(0, 8, f"• {msg}")
                pdf.ln(2)
            pdf.ln(5)

    # Interview Questions Section
    questions = data.get("questionsResults", [])
    if questions:
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(0, 10, "Interview Questions", ln=True)

        # Group by category
        grouped_qs = {}
        for q in questions:
            cat = q.get("category", "General")
            grouped_qs.setdefault(cat, []).append(q.get("question", "").strip())

        for cat, qs in grouped_qs.items():
            if not qs:
                continue
            pdf.set_font("DejaVu", "B", 12)
            pdf.cell(0, 8, cat, ln=True)
            pdf.set_font("DejaVu", "", 11)
            if len(qs) == 1:
                pdf.multi_cell(0, 8, qs[0])
            else:
                for q in qs:
                    pdf.multi_cell(0, 8, f"• {q}")
            pdf.ln(2)
        pdf.ln(5)

    #Improvement Suggestions Section
    mods = data.get("modificationResults", [])
    if mods:
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(0, 10, "Improvement Suggestions", ln=True)

        grouped = {}

        for item in mods:
            msg = item.get("message", "").strip()
            if not msg:
                continue

            # Extract heading before details
            heading_match = re.split(r"(Your CV lists:|JD expects:|Missing:|Suggestion:)", msg, maxsplit=1)
            if len(heading_match) > 1:
                heading = heading_match[0].strip()
                rest = msg[len(heading):].strip()
            else:
                heading = msg.strip()
                rest = ""

            grouped.setdefault(heading, []).append(rest)

        # Print grouped feedback
        for i, (heading, blocks) in enumerate(grouped.items(), 1):
            pdf.set_font("DejaVu", "B", 12)
            pdf.multi_cell(0, 8, f"{heading}")  # heading in bold
            pdf.set_font("DejaVu", "", 12)

            best_block = []
            best_count = -1

            # Special handling for "Missing Core Skills"
            if "Missing Core Skills" in heading:
                for block in blocks:
                    parts = re.split(r"(Your CV lists:|JD expects:|Missing:|Suggestion:)", block)
                    label = None
                    temp_lines = []
                    cv_count = 0

                    for part in parts:
                        if part in ["Your CV lists:", "JD expects:", "Missing:", "Suggestion:"]:
                            label = part
                            continue
                        if label:
                            content = part.strip()
                            if label in ["Your CV lists:", "JD expects:", "Missing:"]:
                                skills = re.split(r'[\s,]+', content)
                                skills = [s.strip("- ") for s in skills if s and s != "-"]
                                if label == "Your CV lists:":
                                    cv_count = len(skills)
                                line = f"{label} {', '.join(skills)}"
                            else:
                                line = f"{label} {content}"
                            temp_lines.append(line)
                            label = None

                    if cv_count > best_count:
                        best_count = cv_count
                        best_block = temp_lines
            else:
                # For all other headings: just take the first block
                block = blocks[0]
                parts = re.split(r"(Your CV lists:|JD expects:|Missing:|Suggestion:)", block)
                label = None
                temp_lines = []
                for part in parts:
                    if part in ["Your CV lists:", "JD expects:", "Missing:", "Suggestion:"]:
                        label = part
                        continue
                    if label:
                        content = part.strip()
                        if label in ["Your CV lists:", "JD expects:", "Missing:"]:
                            skills = re.split(r'[\s,]+', content)
                            skills = [s.strip("- ") for s in skills if s and s != "-"]
                            line = f"{label} {', '.join(skills)}"
                        else:
                            line = f"{label} {content}"
                        temp_lines.append(line)
                        label = None
                best_block = temp_lines

            # Print the chosen block
            for line in best_block:
                pdf.multi_cell(0, 8, line)

            pdf.ln(4)

    #Job Matching Results Section
    matching = data.get("matchingResults", [])
    if matching:
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(0, 10, "Job Matching Results", ln=True)

        # Group by type
        grouped_match = {}
        for m in matching:
            cat = m.get("type", "General")
            grouped_match.setdefault(cat, []).append(m.get("message", "").strip())

        for cat, msgs in grouped_match.items():
            if not msgs:
                continue
            pdf.set_font("DejaVu", "B", 12)
            pdf.cell(0, 8, cat, ln=True)
            pdf.set_font("DejaVu", "", 11)

            if len(msgs) == 1:
                pdf.multi_cell(0, 8, msgs[0])
            else:
                for msg in msgs:
                    pdf.multi_cell(0, 8, f"• {msg}")

            pdf.ln(2)
        pdf.ln(5)

    pdf_bytes = pdf.output(dest="S")
    if isinstance(pdf_bytes, str):  # fpdf v1
        return pdf_bytes.encode("latin-1")
    return pdf_bytes

def rank_feedback(feedback_list):
    priority_map = {
        "missing_metrics": 5,
        "missing_skills": 5,
        "format_issue": 4,
        "outdated_stack": 3,
        "generic_summary": 3,
        "grammar": 2,
        "llm": 3,
        "summary": 3,
        "skills": 4,
        "experience": 5,
        "education": 3,
        "format": 4,
        "improvement": 5,
        "strength": 1,
        "ats optimization": 5
    }

    for fb in feedback_list:
        fb["priority"] = priority_map.get(fb.get("type", "").lower(), 2)

    sorted_feedback = sorted(feedback_list, key=lambda x: -x["priority"])
    return sorted_feedback


# deduplicate_feedback
def deduplicate_feedback(feedback_list):
    if not feedback_list:
        return []

    # Group feedback by intent
    missing_section_feedback = []
    skill_suggestion_feedback = []
    ats_feedback = []
    insights_feedback = []
    other_feedback = []
    seen_messages = set()

    for fb in feedback_list:
        msg_lower = fb["message"].lower().lstrip('• ').strip()
        if msg_lower in seen_messages:
            continue
        seen_messages.add(msg_lower)

        if fb["type"] == "Insights":
            insights_feedback.append(fb)
        elif "missing" in msg_lower or any(f"add a '{sec}' section" in msg_lower for sec in
                                           ["summary", "skills", "experience", "education", "projects",
                                            "certifications"]):
            missing_section_feedback.append(fb["message"])
        elif any(term in msg_lower for term in ["technical skills", "tech stack", "modern tools"]):
            skill_suggestion_feedback.append(fb["message"])
        elif any(term in msg_lower for term in ["ats", "keyword", "formatting", "bullet", "header", "metrics"]):
            cleaned = re.sub(r'^•?\s*ATS Tips:\s*-?', '', fb["message"]).strip()
            ats_feedback.append(cleaned)
        else:
            other_feedback.append(fb)

    # Consolidate missing section feedback
    consolidated = []
    if missing_section_feedback:
        sections_mentioned = set()
        for msg in missing_section_feedback:
            for sec in ["Summary", "Skills", "Experience", "Education", "Projects", "Certifications"]:
                if sec.lower() in msg.lower():
                    sections_mentioned.add(sec)
        consolidated_msg = f"Resume is missing the following sections: {', '.join(sorted(sections_mentioned))}. Consider adding them to improve ATS compatibility."
        consolidated.append({
            "type": "ATS Optimization",
            "message": f" ATS Tips:\n- {consolidated_msg}"
        })

    # Consolidate skill suggestion feedback
    if skill_suggestion_feedback:
        modern_skills = ["react", "aws", "docker", "kubernetes", "tensorflow"]
        detected_skills = []
        for fb in insights_feedback:
            if "detected technical skills" in fb["message"].lower():
                detected_skills = fb["message"].split(": ")[1].split(", ")
        missing_skills = [skill for skill in modern_skills if skill not in detected_skills]
        if missing_skills:
            consolidated.append({
                "type": "Suggestion",
                "message": f"Update your tech stack to include modern tools like {', '.join(missing_skills[:3])}."
            })

    # Add other ATS feedback
    if ats_feedback:
        ats_feedback = [msg for msg in ats_feedback if not any(sec.lower() in msg.lower() for sec in
                                                               ["summary", "skills", "experience", "education",
                                                                "projects", "certifications"])]
        if ats_feedback:
            consolidated.append({
                "type": "ATS Optimization",
                "message": " ATS Tips:\n- " + "\n- ".join(ats_feedback)
            })

    # Add insights and other feedback
    consolidated.extend(insights_feedback)
    consolidated.extend(other_feedback)

    return consolidated
    logger.debug(f"Similarity scores for feedback {i}: {scores[0].tolist()}")

    # Consolidate skill suggestion feedback
    if skill_suggestion_feedback:
        modern_skills = ["react", "aws", "docker", "kubernetes", "tensorflow"]
        detected_skills = []
        for msg in other_feedback:
            if "Detected skills:" in msg["message"]:
                detected_skills = msg["message"].split(": ")[1].split(", ")
        missing_skills = [skill for skill in modern_skills if skill not in detected_skills]
        if missing_skills:
            consolidated.append({
                "type": "Suggestion",
                "message": f"Update your tech stack to include modern tools like {', '.join(missing_skills[:3])}."
            })

    # Add other ATS feedback, excluding missing section messages
    if ats_feedback:
        ats_feedback = [msg for msg in ats_feedback if not any(sec.lower() in msg.lower() for sec in
                                                               ["summary", "skills", "experience", "education",
                                                                "projects", "certifications"])]
        if ats_feedback:
            consolidated.append({
                "type": "ATS Optimization",
                "message": " ATS Tips:\n- " + "\n- ".join(ats_feedback)
            })

    # Add non-ATS feedback
    consolidated.extend(other_feedback)

    return consolidated

def paraphrase_tip(tip):
    prompt = f"Paraphrase this suggestion professionally: {tip}"
    inputs = t5_tokenizer(prompt, return_tensors="pt").to(device)
    outputs = t5_model.generate(inputs['input_ids'], max_length=64)
    return t5_tokenizer.decode(outputs[0], skip_special_tokens=True)


def detect_red_flags(sections):
    flags = []

    experience_raw = sections.get("experience", None)
    experience_text = normalize_section(experience_raw).strip()
    summary_text = normalize_section(sections.get("summary", "")).strip()
    education_text = normalize_section(sections.get("education", "")).strip()

    # --- Experience checks ---
    if experience_raw and experience_text:
        experience_word_count = len(experience_text.split())
        if experience_word_count <= 5:
            flags.append("Experience section is too brief. Add more accomplishments.")
        elif not re.findall(r"[\*\-\u2022]\s", experience_text):
            flags.append("Experience section lacks bullet points.")

    #Education checks
    if not re.search(r"\b\d{4}\b", education_text):
        flags.append("Education section lacks graduation year.")

    # Summary checks
    if re.search(r"\b(hardworking|motivated|team player|passionate)\b", summary_text, re.IGNORECASE):
        flags.append("Summary uses buzzwords. Use results-oriented language instead.")

    return [{"type": "Red Flag", "message": f} for f in flags]

soft_skills = ["communication", "leadership", "problem-solving", "time management", "teamwork", "adaptability"]

def extract_soft_skills(text):
    found = [skill for skill in soft_skills if skill in text.lower()]
    return found

def generate_insights(sections):
    insights = []

    # Check for quantifiable results in projects or experience
    for section_name in ["projects", "experience"]:
        section_text = normalize_section(sections.get(section_name, "")).strip()
        word_count = len(section_text.split())
        if section_text and word_count > 20:  # Require more substantial content
            if re.search(
                r"\d+%|\d+\+?\s?(users|downloads|clients)|reduced\s+.*\d+%|improved\s+.*\d+%|increased\s+.*\d+%",
                section_text, re.IGNORECASE
            ):
                insights.append(f"Demonstrated quantifiable results in the {section_name.title()} section.")

    # If modern technologies are mentioned
    if "skills" in sections:
        skill_text = normalize_section(sections.get("skills", ""))
        skill_list = extract_skills_from_sections(sections)
        if skill_list:
            insights.append(f"Detected technical skills: {', '.join(skill_list)}.")

    # If soft skills are detected
    soft_skills = ["leadership", "communication", "teamwork", "adaptability", "time management"]
    for section_name in ["summary", "experience"]:
        text = normalize_section(sections.get(section_name, "")).lower()
        for soft_skill in soft_skills:
            if soft_skill in text:
                insights.append(f"Highlights soft skill: {soft_skill.title()}.")

    return [{"type": "Insights", "message": f" {msg}"} for msg in insights]

def compute_section_match(resume_sections, job_description):
    """
    Compare each relevant section of the resume to the job description
    and return a match score per section (in percentage form).
    """
    section_scores = {}
    jd_text = clean_text(job_description)

    for section_name, section_content in resume_sections.items():
        section_text = normalize_section(section_content).strip()
        if not section_text:
            continue
        section_text = clean_text(section_text)
        score = compute_similarity(section_text, jd_text)

        if score <= 1:
            score = score * 100
        section_scores[section_name] = round(score, 2)

    if "experience" not in section_scores and "projects" in section_scores:
        section_scores["experience"] = section_scores["projects"]

    return section_scores
